"""
Resume Parser Module
=====================
Extracts raw text from uploaded resume files (PDF and DOCX formats).

Design Decisions:
- pdfplumber is used for PDF parsing because it handles complex layouts
  (tables, columns, headers/footers) far better than PyPDF2.
- python-docx is used for DOCX parsing as it's the standard library for
  reading Word documents in Python.
- Both parsers return clean text with normalized whitespace.
- The parse() function auto-detects file type from extension.

This is the first step in the resume processing pipeline:
    Upload → [Parse] → Preprocess → Extract Features → Classify → Recommend
"""

import os
import re
import pdfplumber
from docx import Document


class ResumeParser:
    """
    Extracts raw text from PDF and DOCX resume files.

    Usage:
        parser = ResumeParser()
        text = parser.parse('/path/to/resume.pdf')
    """

    def parse(self, file_path):
        """
        Parse a resume file and return extracted text.

        Args:
            file_path: Absolute path to the resume file.

        Returns:
            Extracted text as a string, or empty string if extraction fails.

        Raises:
            ValueError: If file type is not supported.
        """
        ext = os.path.splitext(file_path)[1].lower()

        if ext == '.pdf':
            return self._parse_pdf(file_path)
        elif ext == '.docx':
            return self._parse_docx(file_path)
        else:
            raise ValueError(f'Unsupported file type: {ext}. Only PDF and DOCX are supported.')

    def _parse_pdf(self, file_path):
        """
        Extract text from a PDF file using pdfplumber.

        pdfplumber extracts text page-by-page, preserving the reading order.
        It handles multi-column layouts better than most alternatives.

        Args:
            file_path: Path to the PDF file.

        Returns:
            Extracted text string.
        """
        text_parts = []

        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
        except Exception as e:
            print(f'Error parsing PDF {file_path}: {e}')
            return ''

        raw_text = '\n'.join(text_parts)
        return self._normalize_whitespace(raw_text)

    def _parse_docx(self, file_path):
        """
        Extract text from a DOCX file using python-docx.

        Extracts text from:
        1. All paragraphs (body text, headings)
        2. All tables (structured data like education, experience)

        Args:
            file_path: Path to the DOCX file.

        Returns:
            Extracted text string.
        """
        text_parts = []

        try:
            doc = Document(file_path)

            # Extract paragraph text
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())

            # Extract table text (resumes often use tables for layout)
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(' | '.join(row_text))

        except Exception as e:
            print(f'Error parsing DOCX {file_path}: {e}')
            return ''

        raw_text = '\n'.join(text_parts)
        return self._normalize_whitespace(raw_text)

    @staticmethod
    def _normalize_whitespace(text):
        """
        Normalize whitespace in extracted text.

        - Replace multiple spaces with single space
        - Replace multiple newlines with double newline (paragraph separator)
        - Strip leading/trailing whitespace
        """
        # Replace multiple spaces with single space
        text = re.sub(r'[ \t]+', ' ', text)
        # Replace 3+ newlines with double newline
        text = re.sub(r'\n{3,}', '\n\n', text)
        # Strip each line
        lines = [line.strip() for line in text.split('\n')]
        return '\n'.join(lines).strip()


# Singleton instance for convenience
resume_parser = ResumeParser()
